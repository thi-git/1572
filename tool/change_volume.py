# -*- coding: utf-8 -*-
import os
import json
import pandas as pd
import shutil
from config import DevelopmentConfig as CFG
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.ext.declarative import declarative_base
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

Base = declarative_base()
engine = create_engine(CFG.SQLALCHEMY_DATABASE_URI, echo=True)
DB_session = sessionmaker(engine)
db_session = DB_session()


def wash():
    # 最新模板存放在download_temp
    # 需要清洗的檔案存放在old_file
    # Step1. 取得所有待清洗的檔案(convert不處理/鼎漢的做清洗/上傳檔的複製原檔)
    # Step2. 將前兩頁內容暫存為json+存取第三頁的PCE資料
    # Step3. 取得對應的下載模板
    # Step4. 產出新的下載檔並儲存到export(將第一二頁全部內容+第三頁PCE資料存到新檔案)

    all_folders = os.listdir(CFG.OLD_FOLDER)
    folder_idx = 1
    for each_folder in all_folders:
        # 1. 取得所有待清洗的檔案
        all_files = os.listdir(f"{CFG.OLD_FOLDER}/{each_folder}")
        for file in all_files:
            # convert檔案不處理
            if 'convert' not in file:
                # 鼎漢=>更新檔案/上傳檔=>複製原檔
                if '鼎漢' in file:
                    sheet_list = ['轉向量', '路口基本資料']
                    file_path = os.path.join(f"{CFG.OLD_FOLDER}/{each_folder}", file)

                    # 2. 將前兩頁內容暫存為json+存取第三頁的PCE資料
                    data_list = {}
                    for sheet in sheet_list:
                        df = pd.read_excel(file_path, sheet_name=sheet)
                        data_list[sheet] = df.to_json(force_ascii=False)

                    wb_org = load_workbook(file_path)

                    # 取得分頁二部分資料
                    ws_p2 = wb_org['路口基本資料']
                    tc_id = ws_p2['B4'].value
                    date = str(ws_p2['B8'].value).replace('/', '-')
                    is_holiday = '平日' if ws_p2['B9'].value == '平日' else '假日'
                    intersection_type = ws_p2['B13'].value

                    # 取得分頁三PCE資料
                    ws_p3 = wb_org['調查資料(OUT)']
                    pce_data = []
                    for cell in ws_p3[1]:
                        if cell.value is None:  # 如果儲存格為空
                            break
                        pce_data.append(cell.value)

                    # 3. 取得對應的下載模板
                    fn = f"{CFG.DOWNLOAD_FOLDER}/volume/{intersection_type}_{is_holiday}.xlsx"
                    wb = load_workbook(fn)
                    sheet_list = ['路口基本資料', '轉向量']

                    # 4. 產出新的下載檔並儲存到export
                    for sheet in sheet_list:
                        ws = wb[sheet]

                        # 將json檔轉為dataframe
                        df_page = pd.DataFrame.from_dict(json.loads(data_list[sheet]))

                        if sheet == '轉向量':
                            header = True
                        elif sheet == '路口基本資料':
                            header = False

                            # 依照不同路型調整資料欄數
                            if intersection_type == '正交四叉路口':
                                data_num = 6
                            elif intersection_type == '五叉路口':
                                data_num = 7
                            elif intersection_type == '六叉路口':
                                data_num = 8

                            df_page = df_page.iloc[:, 0:data_num]
                            new_df = pd.DataFrame([df_page.columns], columns=df_page.columns)
                            new_df = new_df.iloc[:, 0:data_num]
                            for col in df_page.columns[2:]:
                                if 'Unnamed' in col:
                                    new_df[col] = ''
                            df_page = pd.concat([new_df, df_page], ignore_index=True)

                        # 將dataframe存入excel
                        rows = dataframe_to_rows(df_page, index=False, header=header)
                        for r_idx, row in enumerate(rows, 1):
                            for c_idx, value in enumerate(row, 1):
                                ws.cell(row=r_idx, column=c_idx, value=value)

                    # 第三頁資料
                    ws_new = wb['調查資料(OUT)']
                    for col_num, value in enumerate(pce_data, start=1):
                        ws_new.cell(row=1, column=col_num, value=value)

                    # 存檔
                    save_path = f"export/volume/{tc_id}"
                    file_name = f"{tc_id}_{date}_volume_鼎漢.xlsx"

                    # 檢查目錄是否存在，若不存在則創建目錄
                    if not os.path.isdir(f"{CFG.STATIC_FOLDER}/{save_path}"):
                        os.makedirs(f"{CFG.STATIC_FOLDER}/{save_path}")

                    export_path = f"{CFG.STATIC_FOLDER}/{save_path}/{file_name}"
                    wb.save(export_path)
                    print(f"編號: {folder_idx}, TC編號: {tc_id}, 檔案路徑: {export_path}")
                if '上傳檔' in file:
                    # 舊檔案路徑
                    org_file_path = os.path.join(f"{CFG.OLD_FOLDER}/{each_folder}", file)

                    # 新檔案路徑
                    destination_file = f"{CFG.EXPORT_FOLDER}/volume/{each_folder}/{file}"

                    # 檢查目錄是否存在，若不存在則創建目錄
                    if not os.path.isdir(f"{CFG.EXPORT_FOLDER}/volume/{each_folder}"):
                        os.makedirs(f"{CFG.EXPORT_FOLDER}/volume/{each_folder}")

                    # 複製檔案並儲存
                    shutil.copy(org_file_path, destination_file)
                    print(f"編號: {folder_idx}, TC編號: {each_folder}, 檔案路徑: {destination_file}")
        folder_idx = folder_idx + 1


wash()

